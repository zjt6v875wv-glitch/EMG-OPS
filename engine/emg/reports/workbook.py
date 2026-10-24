"""Master payroll workbook (values, with live total formulas) and CSV exports."""
from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..money import D, cents
from .data import ReportData

HEAD = PatternFill("solid", fgColor="173247")
HFONT = Font(color="FFFFFF", bold=True)
MONEY = '#,##0.00;[Red]-#,##0.00'

PEOPLE_COLS = [
    ("payee", "Person"), ("status", "Status"), ("agent_commission", "Agent pay"), ("manager_override", "Overrides / premium"),
    ("gross_compensation", "Gross"), ("opening_chargeback", "Opening CB"), ("new_chargebacks", "New CB"), ("held_chargebacks", "Held CB (unverified)"),
    ("chargeback_recovery", "CB recovered"), ("ending_chargeback", "Ending CB"), ("manual_hold", "Held pay"), ("reserve_withheld", "Reserve"),
    ("room_opening", "Room opening"), ("room_new", "Room new"), ("room_deduction", "Room collected"), ("room_ending", "Room remaining"),
    ("advance_opening", "Advance opening"), ("new_advance", "New advance"), ("advance_deduction", "Advance collected"), ("advance_ending", "Advance remaining"),
    ("flight_opening", "Flight opening"), ("flight_new", "Flight new"), ("flight_deduction", "Flight collected"), ("flight_ending", "Flight remaining"),
    ("other_opening", "Other opening"), ("other_new", "Other new"), ("other_deduction", "Other collected"), ("other_ending", "Other remaining"),
    ("protected_check", "Protected check"), ("commission_payment", "Commission payment"), ("cash_credits", "Cash credits"), ("final_check", "Final check"),
]


def _num(v):
    if isinstance(v, Decimal):
        return float(v)  # openpyxl cell value; workbook display only, source of truth stays in the database
    if isinstance(v, str):
        try:
            return float(Decimal(v))
        except Exception:
            return v
    return v


def _sheet(wb, title, header, rows, money_cols=(), totals=False, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD, HFONT, Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([_num(v) if i in money_cols else (v.isoformat() if isinstance(v, date) else v) for i, v in enumerate(r)])
    n = len(rows) + 1
    for i in money_cols:
        col = get_column_letter(i + 1)
        for row in ws.iter_rows(min_row=2, max_row=n, min_col=i + 1, max_col=i + 1):
            row[0].number_format = MONEY
        if totals and n > 1:
            ws[f"{col}{n + 1}"] = f"=ROUND(SUM({col}2:{col}{n}),2)"
            ws[f"{col}{n + 1}"].number_format = MONEY
            ws[f"{col}{n + 1}"].font = Font(bold=True)
    if totals and n > 1:
        ws[f"A{n + 1}"] = "TOTAL"
        ws[f"A{n + 1}"].font = Font(bold=True)
    ws.freeze_panes = "B2"
    for i, h in enumerate(header):
        ws.column_dimensions[get_column_letter(i + 1)].width = (widths or {}).get(i, max(10, min(40, len(str(h)) + 4)))
    return ws


def master_workbook(rd: ReportData) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    t = rd.run["totals"]
    info = [("Payroll week", rd.week), ("Pay date (Friday)", rd.pay_date.isoformat()), ("Status", rd.label), ("Revision", rd.run["revision"]),
            ("Rule set version", rd.run["rule_set_version"]), ("Finalized by", rd.run.get("finalized_by") or ""), ("Correction reason", rd.run.get("correction_reason") or ""),
            ("Total checks", t.get("total_checks")), ("Payable checks", t.get("positive_checks")), ("Gross compensation", t.get("gross_compensation")),
            ("Chargeback recovery", t.get("chargeback_recovery")), ("Held pay", t.get("manual_holds")), ("Room deductions", t.get("room_deductions")),
            ("Advance deductions", t.get("advance_deductions")), ("Flight deductions", t.get("flight_deductions")), ("Other deductions", t.get("other_deductions")),
            ("Cash credits", t.get("cash_credits"))] + [(f"{k} (company payee, excluded from payroll)", v) for k, v in (t.get("company_payables") or {}).items()]
    ws.append(["EMG payroll", ""])
    ws["A1"].font = Font(bold=True, size=14, color="173247")
    for k, v in info:
        ws.append([k, _num(v) if isinstance(v, str) and k not in ("Payroll week", "Pay date (Friday)", "Status", "Finalized by", "Correction reason") else v])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(row[0].value, float):
            row[0].number_format = MONEY
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 48, 24

    money_idx = set(range(2, len(PEOPLE_COLS)))
    _sheet(wb, "Payroll Calculation", [h for _, h in PEOPLE_COLS], [[p["data"].get(k, "") for k, _ in PEOPLE_COLS] for p in rd.people], money_idx, totals=True, widths={0: 28})
    pay = [p for p in rd.people if D(p["final_check"]) > 0]
    _sheet(wb, "Payable Checks", ["Person", "Final check", "Advance remaining", "Advance collected", "Room collected", "Flight collected", "Cash credits"],
           [[p["payee"]] + [p["data"][k] for k in ("final_check", "advance_ending", "advance_deduction", "room_deduction", "flight_deduction", "cash_credits")]
            for p in sorted(pay, key=lambda p: -D(p["final_check"]))], set(range(1, 7)), totals=True, widths={0: 28})
    _sheet(wb, "Earnings Detail", ["Payee", "Activity", "Vendor", "Selling agent", "Customer", "Utility", "Product", "kWh", "Units", "Rate", "Pay", "Enrollment", "Transaction", "Sale date", "Rule", "Notes"],
           [[l["payee"], l["activity"], l["vendor"], l["rep"], l["customer"], l["utility"], l["product"], _num(l["kwh"]), _num(l["units"]), _num(l["rate"]), l["amount"],
             l["enrollment_id"], l["transaction_id"], l["sale_date"], l["rule_id"], l["notes"]] for l in rd.lines], {10}, totals=True)
    _sheet(wb, "Chargebacks", ["Payee", "Vendor", "Enrollment", "Selling agent", "Customer", "Reversal", "Verified", "Original week", "Evidence"],
           [[c["payee"], c["vendor"], c["enrollment_id"], c["data"].get("rep"), c["data"].get("customer"), c["amount"], "yes" if c["verified"] else "HELD",
             c["data"].get("original_week"), c["data"].get("evidence")] for c in rd.chargebacks], {5}, totals=True, widths={8: 60})
    comp = [i for i in rd.items if i["kind"] == "company_adjustment"]
    _sheet(wb, "Company Adjustments", ["Vendor", "Enrollment", "Customer", "Amount", "Treatment"],
           [[i["vendor"], i["enrollment_id"], i["data"].get("customer"), i["data"].get("amount"), i["data"].get("reason")] for i in comp], {3}, totals=True, widths={4: 70})
    for kind, title in (("nonpayable", "Non-payable"), ("pending", "Pending")):
        _sheet(wb, title, ["Agent / company", "Vendor", "Customer", "Utility", "Product", "Enrollment", "Sale date", "Status"],
               [[i["payee"], i["vendor"], i["data"].get("customer"), i["data"].get("utility"), i["data"].get("product"), i["enrollment_id"], i["data"].get("sale_date"), i["data"].get("status")]
                for i in rd.items if i["kind"] == kind], widths={7: 60})
    _sheet(wb, "Week Adjustments", ["Person", "Kind", "Amount", "Method", "Value", "Debts", "Creditor", "Label", "Note", "Authorized by", "Reference"],
           [[a["person"], a["kind"], a["amount"], a["method"], _num(a["value"]), ", ".join(a["debts"] or []), a["creditor"], a["label"], a["note"], a["authorized_by"], a["external_ref"]] for a in rd.adjustments], {2})
    _sheet(wb, "Review", ["Blocking", "Status", "Kind", "Message", "Vendor", "Enrollment", "Person", "Amount", "Resolution"],
           [["yes" if r["blocking"] else "", r["status"], r["kind"], r["message"], r["vendor"], r["enrollment_id"], r["person"], r["amount"], r["resolution"]] for r in rd.review], {7}, widths={3: 70, 8: 50})
    _sheet(wb, "Vendor Reconciliation", ["Vendor", "Vendor net (after chargebacks)", "Payroll cost", "Company retained"],
           [[v, x["vendor_revenue"], x["payroll_cost"], x["company_retained"]] for v, x in (rd.run.get("reconciliation") or {}).items()], {1, 2, 3}, totals=True)
    _sheet(wb, "Sources", ["Importer", "Vendor", "File", "SHA-256", "Rows"],
           [[s["importer"], s["vendor"], s["file_name"], s["sha256"], ", ".join(f"{v} {k}" for k, v in (s["summary"].get("counts") or {}).items())] for s in rd.sources], widths={2: 50, 3: 66, 4: 50})
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _csv(header, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8")


def payable_csv(rd: ReportData) -> bytes:
    pay = sorted((p for p in rd.people if D(p["final_check"]) > 0), key=lambda p: -D(p["final_check"]))
    return _csv(["Agent", "Final Check", "Advances Owed", "Advance Deduction", "Room Deduction", "Flight Deduction", "Cash Credits"],
                [[p["payee"]] + [f"{cents(D(p['data'][k])):.2f}" for k in ("final_check", "advance_ending", "advance_deduction", "room_deduction", "flight_deduction", "cash_credits")] for p in pay])


def payable_only_csv(rd: ReportData) -> bytes:
    pay = sorted((p for p in rd.people if D(p["final_check"]) > 0), key=lambda p: p["payee"].upper())
    return _csv(["Agent", "Final Check"], [[p["payee"], f"{cents(D(p['final_check'])):.2f}"] for p in pay])


def paychex_csv(rd: ReportData, paychex_names: dict[str, str] | None = None) -> bytes:
    """Paychex contractor import: one row per positive check, amount in the 1099-NEC Amount field."""
    names = paychex_names or {}
    pay = sorted((p for p in rd.people if D(p["final_check"]) > 0), key=lambda p: p["payee"].upper())
    return _csv(["Contractor", "1099-NEC Amount"], [[names.get(p["payee"], p["payee"]), f"{cents(D(p['final_check'])):.2f}"] for p in pay])


def company_invoice_csv(rd: ReportData, payee: str, invoice_no: str, source_invoice: str = "") -> bytes:
    rows = sorted((l for l in rd.lines if l["payee"] == payee), key=lambda l: (l["rep"], str(l["sale_date"]), l["enrollment_id"]))
    return _csv(["Invoice", "Source Invoice", "Agent", "Enrollment ID", "Transaction ID", "Sold Date", "Customer", "Enrollment Count", "Rate", "Amount"],
                [[invoice_no, source_invoice, l["rep"], l["enrollment_id"], l["transaction_id"], l["sale_date"] or "", l["customer"], 1, f"{D(l['rate']):.2f}", f"{cents(D(l['amount'])):.2f}"] for l in rows])
