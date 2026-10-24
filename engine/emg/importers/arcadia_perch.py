"""Arcadia / Perch "Dealer Commission Statement" workbook.

Sheets: Payable (paid this statement), Future Payable (pending),
Lookback_Non_Payable (denied / cannot process / in progress).
Campaign is normalized by product family: DUAL* -> Arcadia, UCB* -> Perch.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from ..model import SaleRow
from ..names import clean, key, ref_key
from .base import ImportResult, FieldMapping, col, find_header, num, sha256_file, to_date

COLUMNS = ["Campaign", "Type", "Utility Name", "Rep Name", "Customer Name", "Phone", "Sale Date", "Status", "Payable kWh", "Commissions", "Reference #", "Product"]
COMPANY_ONLY_STAGE = ["ADJUSTMENT", "SECOND", "FINAL", "REMAINING", "REMAINDER", "BACKEND", "BACK END"]


def detect(path: str | Path) -> bool:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    return "Payable" in wb.sheetnames and any("Commission Statement" in str(c) for r in wb["Payable"].iter_rows(max_row=3, values_only=True) for c in r if c)


def normalized_vendor(raw: str, product: str) -> str:
    p = key(product)
    if p.startswith("DUAL"):
        return "Arcadia"
    if p.startswith("UCB"):
        return "Perch"
    return clean(raw)


def commercial_stage(status: str) -> str:
    s = key(status)
    if any(k in s for k in COMPANY_ONLY_STAGE):
        return "COMPANY ONLY"
    if "UPFRONT" in s:
        return "INITIAL"
    return "UNKNOWN"


def parse(path: str | Path, year: int | None = None) -> ImportResult:
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    res = ImportResult("Arcadia/Perch", "dealer_commission_statement", path.name, sha256_file(path))
    pay = list(wb["Payable"].iter_rows(values_only=True))
    ach = None
    total = None
    for r in pay[:8]:
        for j, c in enumerate(r):
            if key(c) == "ACH DATE" and j + 1 < len(r):
                ach = to_date(r[j + 1])
            if key(c) == "PAYOUT" and j + 2 < len(r):
                total = r[j + 2]
    res.meta.update(invoice_date=ach, vendor_total=str(num(total)) if total is not None else None)
    hi, idx = find_header(pay, ["Campaign", "Rep Name", "Reference #"])
    res.mapping = [FieldMapping(t, c, t in ("Rep Name", "Reference #"), pay[hi + 1][idx[key(c)]] if len(pay) > hi + 1 and key(c) in idx else None)
                   for t, c in [(c, c) for c in COLUMNS]]
    invoice_number = f"ARC-{ach.isoformat()}" if ach else path.stem

    def make(r, i, sheet, line_type):
        raw_vendor = clean(col(r, idx_s, "Campaign"))
        product = clean(col(r, idx_s, "Product"))
        status = clean(col(r, idx_s, "Status"))
        premise = clean(col(r, idx_s, "Type")) or "Residential"
        stage = ""
        if key(premise) == "COMMERCIAL":
            product = "COMMERCIAL"
            stage = commercial_stage(status)
        if not product and "0.015 UPFRONT" in key(status):
            product = "UCB"
        amount = num(col(r, idx_s, "Commissions") if col(r, idx_s, "Commissions") is not None else col(r, idx_s, "Comms"))
        return SaleRow(
            vendor=normalized_vendor(raw_vendor, product), raw_vendor=raw_vendor,
            enrollment_id=ref_key(col(r, idx_s, "Reference #")), raw_rep=clean(col(r, idx_s, "Rep Name")),
            customer=clean(col(r, idx_s, "Customer Name")), phone=clean(col(r, idx_s, "Phone")),
            utility=clean(col(r, idx_s, "Utility Name")), product=product, premise=premise,
            kwh=num(col(r, idx_s, "Payable kWh")), sale_date=to_date(col(r, idx_s, "Sale Date")),
            invoice_date=ach if sheet == "Payable" else None, invoice_number=invoice_number if sheet == "Payable" else "",
            vendor_amount=amount, line_type=line_type, status=status, stage=stage,
            source_file=path.name, source_sheet=sheet, source_row=i + 1,
        )

    idx_s = idx
    for i in range(hi + 1, len(pay)):
        r = pay[i]
        if not col(r, idx, "Campaign") or not col(r, idx, "Rep Name"):
            continue
        s = make(r, i, "Payable", "payable")
        if s.vendor_amount < 0 or "CHARGEBACK" in key(s.status):
            s.line_type = "chargeback"
        elif "ADJUSTMENT" in key(s.status):
            s.line_type = "adjustment"
        elif s.stage == "COMPANY ONLY":
            s.line_type = "company_only"
            s.notes = "Commercial later payment: entirely EMG; zero agent and manager pay"
        elif s.stage == "UNKNOWN":
            s.line_type = "review"
            s.notes = "Commercial payment stage requires review; do not pay a second installment"
        if s.line_type == "payable" and year and s.sale_date and s.sale_date.year != year:
            s.line_type = "review"
            s.notes = f"Sale date {s.sale_date} is outside {year}; confirm before paying"
        if not s.enrollment_id:
            res.errors.append({"sheet": "Payable", "row": i + 1, "error": "Missing Reference #"})
        res.rows.append(s)

    for sheet, line_type in (("Future Payable", "pending"), ("Lookback_Non_Payable", "nonpayable")):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        try:
            h2, idx_s = find_header(rows, ["Campaign", "Rep Name", "Reference #"])
        except ValueError:
            res.warnings.append({"sheet": sheet, "warning": "header not found"})
            continue
        for i in range(h2 + 1, len(rows)):
            r = rows[i]
            if not col(r, idx_s, "Rep Name"):
                continue
            s = make(r, i, sheet, line_type)
            if line_type == "nonpayable" and "PENDING" in key(s.status):
                s.line_type = "pending"
            if year and s.sale_date and s.sale_date.year != year:
                res.excluded.append({"sheet": sheet, "row": i + 1, "reference": s.enrollment_id, "reason": f"prior-year sale {s.sale_date}"})
                continue
            res.rows.append(s)
    return res
